# Tests for utils/data_pipeline.py.

import uuid
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest

import utils.data_pipeline as dp


class TestChunkedDocuments:

    def test_text_structure_based_returns_chunks(self):
        text = "word " * 500  # 2500 chars -> at least 3 chunks of 1000
        result = dp.chunked_documents(text, chunking_method="text_structure_based")
        assert isinstance(result, list)
        assert len(result) >= 2
        assert all(isinstance(chunk, str) for chunk in result)

    def test_text_structure_based_is_default(self):
        text = "some content here"
        assert (dp.chunked_documents(text) ==
                dp.chunked_documents(text, chunking_method="text_structure_based"))

    def test_empty_text_returns_empty_list(self):
        assert dp.chunked_documents("", chunking_method="text_structure_based") == []

    def test_markdown_based_splits_on_headers(self):
        text = "# Title\n\nintro text\n\n## Section\n\nsection text"
        result = dp.chunked_documents(text, chunking_method="markdown_based")
        assert len(result) >= 1
        # Header content is split into separate documents; the header text ends
        # up in metadata, the body in page_content.
        assert any("section text" in getattr(doc, "page_content", "") or "section text" in doc
                   for doc in result)

    def test_markdown_based_without_headers_returns_single_chunk(self):
        text = "just plain text without any headers"
        result = dp.chunked_documents(text, chunking_method="markdown_based")
        assert len(result) == 1

    def test_semantic_chunker_uses_embedder(self):
        expected = ["chunk one", "chunk two"]
        with patch.object(dp, "SemanticChunker") as mock_splitter_cls:
            mock_splitter_cls.return_value.create_documents.return_value = expected
            result = dp.chunked_documents("some text", chunking_method="semantic_chunker")

        assert result == expected
        mock_splitter_cls.assert_called_once()
        mock_splitter_cls.return_value.create_documents.assert_called_once_with(["some text"])


class TestStoreInVectorDb:

    def test_creates_collection_and_returns_uuid_ids(self, tmp_path):
        collection, ids = dp.store_in_vector_db(
            documents=["doc one", "doc two"],
            model_name="test_collection",
            collection_path=str(tmp_path / "chroma"),
        )

        assert len(ids) == 2
        # Every id is a valid uuid string.
        for id_ in ids:
            uuid.UUID(id_)
        assert collection.count() == 2

    def test_persists_across_clients(self, tmp_path):
        collection_path = str(tmp_path / "chroma")
        dp.store_in_vector_db(
            documents=["doc"], model_name="test_collection", collection_path=collection_path
        )
        import chromadb
        client = chromadb.PersistentClient(path=collection_path)
        assert client.get_or_create_collection("test_collection").count() == 1

    def test_get_or_create_reuses_existing_collection(self, tmp_path):
        collection_path = str(tmp_path / "chroma")
        col1, _ = dp.store_in_vector_db(
            documents=["a"], model_name="same_name", collection_path=collection_path
        )
        col2, _ = dp.store_in_vector_db(
            documents=["b"], model_name="same_name", collection_path=collection_path
        )
        assert col2.count() == 2

    def test_custom_embeddings_are_used(self, tmp_path):
        embedding = [0.1] * 768
        collection, ids = dp.store_in_vector_db(
            documents=["image path"],
            model_name="img_collection",
            collection_path=str(tmp_path / "chroma"),
            embeddings=[embedding],
        )
        assert collection.count() == 1
        # The stored vector matches the one we passed in.
        stored = collection.get(ids=ids, include=["embeddings"])
        assert stored["embeddings"][0].tolist() == pytest.approx(embedding)

    def test_empty_documents_list_raises(self, tmp_path):
        # chroma refuses an empty add (documents and ids must be non-empty).
        with pytest.raises(ValueError, match="Non-empty lists are required"):
            dp.store_in_vector_db(
                documents=[], model_name="empty_collection", collection_path=str(tmp_path / "chroma")
            )

    def test_empty_documents_with_embeddings_length_mismatch(self, tmp_path):
        # chroma raises when documents and embeddings lengths disagree.
        with pytest.raises(Exception):
            dp.store_in_vector_db(
                documents=["a", "b"],
                model_name="bad_collection",
                collection_path=str(tmp_path / "chroma"),
                embeddings=[[0.1] * 768],
            )


class TestGetImageEmbedding:

    def _mock_result(self, values):
        result = MagicMock()
        result.embeddings = [SimpleNamespace(values=values)]
        return result

    def test_returns_embedding_values(self):
        with patch.object(dp.client.models, "embed_content",
                          return_value=self._mock_result([0.5, 0.25])) as mock_embed:
            vector = dp.get_image_embedding(b"fake-bytes", mime_type="image/png")

        assert vector == [0.5, 0.25]
        _, kwargs = mock_embed.call_args
        assert kwargs["model"] == dp.GEMINI_EMBEDDING_MODEL
        assert kwargs["contents"][0].inline_data.mime_type == "image/png"

    def test_default_mime_type_is_png(self):
        with patch.object(dp.client.models, "embed_content",
                          return_value=self._mock_result([0.1])) as mock_embed:
            dp.get_image_embedding(b"fake-bytes")
        _, kwargs = mock_embed.call_args
        assert kwargs["contents"][0].inline_data.mime_type == "image/png"



class TestStoreImage:

    def test_none_path_returns_none(self):
        assert dp.store_image(path=None) is None

    def test_stores_png_image(self, tmp_path):
        image_path = tmp_path / "pic.png"
        image_path.write_bytes(b"png-bytes")

        with patch.object(dp, "get_image_embedding", return_value=[0.1] * 768) as mock_embed, \
             patch.object(dp, "store_in_vector_db") as mock_store:
            mock_store.return_value = ("collection", ["id-1"])
            result = dp.store_image(path=str(image_path), model_name="my_model")

        mock_embed.assert_called_once_with(b"png-bytes", mime_type="image/png")
        mock_store.assert_called_once()
        call_kwargs = mock_store.call_args.kwargs
        assert call_kwargs["documents"] == [str(image_path)]
        assert call_kwargs["model_name"] == "my_model"
        assert call_kwargs["embeddings"] == [[0.1] * 768]
        assert result == ("collection", ["id-1"])

    def test_jpg_defaults_to_jpeg_mime(self, tmp_path):
        image_path = tmp_path / "pic.jpg"
        image_path.write_bytes(b"jpg-bytes")

        with patch.object(dp, "get_image_embedding", return_value=[0.2]) as mock_embed, \
             patch.object(dp, "store_in_vector_db"):
            dp.store_image(path=str(image_path), model_name="m")

        mock_embed.assert_called_once_with(b"jpg-bytes", mime_type="image/jpeg")

    def test_non_png_extension_still_treated_as_jpeg(self, tmp_path):
        # Anything not ending in .png (e.g. .gif, .bmp) falls into the jpeg branch.
        image_path = tmp_path / "pic.bmp"
        image_path.write_bytes(b"bmp-bytes")

        with patch.object(dp, "get_image_embedding", return_value=[0.3]) as mock_embed, \
             patch.object(dp, "store_in_vector_db"):
            dp.store_image(path=str(image_path), model_name="m")

        mock_embed.assert_called_once_with(b"bmp-bytes", mime_type="image/jpeg")

    def test_generates_model_name_when_missing(self, tmp_path):
        image_path = tmp_path / "pic.png"
        image_path.write_bytes(b"png-bytes")

        with patch.object(dp, "get_image_embedding", return_value=[0.1]), \
             patch.object(dp, "store_in_vector_db") as mock_store:
            dp.store_image(path=str(image_path))

        model_name = mock_store.call_args.kwargs["model_name"]
        uuid.UUID(model_name)  # raises if not a valid uuid

    def test_missing_file_raises(self):
        with patch.object(dp, "get_image_embedding"), patch.object(dp, "store_in_vector_db"):
            with pytest.raises(FileNotFoundError):
                dp.store_image(path="does/not/exist.png", model_name="m")

    def test_end_to_end_with_real_chroma(self, tmp_path):
        image_path = tmp_path / "pic.png"
        image_path.write_bytes(b"png-bytes")

        # store_image does not expose collection_path, so route the storage call
        # to the real store_in_vector_db in a temp directory.
        real_store = dp.store_in_vector_db
        with patch.object(dp, "get_image_embedding", return_value=[0.4] * 768), \
             patch.object(dp, "store_in_vector_db",
                          side_effect=lambda **kw: real_store(
                              **{**kw, "collection_path": str(tmp_path / "chroma")})):
            collection, ids = dp.store_image(path=str(image_path), model_name="e2e_model")

        assert collection.count() == 1
        assert len(ids) == 1


class TestStorePdf:

    def _make_reader(self, page_texts):
        reader = MagicMock()
        reader.pages = [MagicMock(extract_text=MagicMock(return_value=t)) for t in page_texts]
        return reader

    def test_extracts_text_from_all_pages(self, tmp_path):
        pdf_path = tmp_path / "doc.pdf"
        pdf_path.write_bytes(b"pdf-bytes")

        with patch.object(dp, "PdfReader", return_value=self._make_reader(["page one ", "page two"])), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]) as mock_chunk, \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])) as mock_store:
            result = dp.store_pdf(path=str(pdf_path), model_name="pdf_model")

        mock_chunk.assert_called_once_with("page one page two", chunking_method="text_structure_based")
        mock_store.assert_called_once_with(documents=["chunk"], model_name="pdf_model")
        assert "completed" in result

    def test_custom_chunking_method_is_forwarded(self, tmp_path):
        pdf_path = tmp_path / "doc.pdf"
        pdf_path.write_bytes(b"pdf-bytes")

        with patch.object(dp, "PdfReader", return_value=self._make_reader(["text"])), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]) as mock_chunk, \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])):
            dp.store_pdf(path=str(pdf_path), model_name="m", chunking_method="semantic_chunker")

        assert mock_chunk.call_args.kwargs["chunking_method"] == "semantic_chunker"

    def test_empty_pdf_text(self, tmp_path):
        pdf_path = tmp_path / "empty.pdf"
        pdf_path.write_bytes(b"pdf-bytes")

        with patch.object(dp, "PdfReader", return_value=self._make_reader([])), \
             patch.object(dp, "chunked_documents", return_value=[]) as mock_chunk, \
             patch.object(dp, "store_in_vector_db", return_value=(None, [])):
            result = dp.store_pdf(path=str(pdf_path), model_name="m")

        # Empty list is not None, so the "chunking issue" early return is skipped.
        mock_chunk.assert_called_once_with("", chunking_method="text_structure_based")

    def test_chunking_failure_skips_storage(self, tmp_path):
        pdf_path = tmp_path / "doc.pdf"
        pdf_path.write_bytes(b"pdf-bytes")

        with patch.object(dp, "PdfReader", return_value=self._make_reader(["text"])), \
             patch.object(dp, "chunked_documents", return_value=None), \
             patch.object(dp, "store_in_vector_db") as mock_store:
            result = dp.store_pdf(path=str(pdf_path), model_name="m")

        mock_store.assert_not_called()
        assert result is None

    def test_default_model_name_is_uuid(self, tmp_path):
        pdf_path = tmp_path / "doc.pdf"
        pdf_path.write_bytes(b"pdf-bytes")

        with patch.object(dp, "PdfReader", return_value=self._make_reader(["text"])), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]), \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])) as mock_store:
            dp.store_pdf(path=str(pdf_path))

        uuid.UUID(mock_store.call_args.kwargs["model_name"])


def _tiny_png_bytes():
    from PIL import Image
    buffer = BytesIO()
    Image.new("RGB", (4, 4), color="red").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def pptx_path(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()

    # Slide 1: title + body text.
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Slide One Title"
    slide1.placeholders[1].text = "Body text for slide one."

    # Slide 2: blank layout with a table, a group containing a textbox, and a picture.
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "value 1"
    table.cell(1, 1).text = "  "  # whitespace-only cell -> skipped

    group = slide2.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Grouped text"

    slide2.shapes.add_picture(BytesIO(_tiny_png_bytes()), Inches(6), Inches(1), Inches(2), Inches(2))

    # Speaker notes on slide 2.
    slide2.notes_slide.notes_text_frame.text = "Speaker notes here."

    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return str(path)


class TestExtractPptContent:

    def test_extracts_text_tables_notes_and_images(self, pptx_path):
        text_content, images = dp.extract_ppt_content(pptx_path)

        assert "Slide One Title" in text_content
        assert "Body text for slide one." in text_content
        assert "Grouped text" in text_content              # from inside the group
        assert "Header A | Header B" in text_content       # table row joined with " | "
        assert "value 1" in text_content
        assert "Speaker notes here." in text_content
        assert "  " not in [part for part in text_content.split("\n\n")]  # blank cell dropped

        assert len(images) == 1
        assert images[0]["ext"] == "png"
        assert images[0]["mime_type"] == "image/png"
        assert images[0]["bytes"].startswith(b"\x89PNG")

    def test_ppt_without_images_returns_empty_image_list(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Only text"
        path = tmp_path / "text_only.pptx"
        prs.save(str(path))

        text_content, images = dp.extract_ppt_content(str(path))
        assert "Only text" in text_content
        assert images == []

    def test_ppt_without_notes(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "No notes slide"
        path = tmp_path / "no_notes.pptx"
        prs.save(str(path))

        text_content, _ = dp.extract_ppt_content(str(path))
        assert "No notes slide" in text_content


class TestStorePpt:

    def _fake_image(self):
        return {"bytes": b"png-bytes", "mime_type": "image/png", "ext": "png"}

    @pytest.fixture(autouse=True)
    def _run_in_tmp_dir(self, tmp_path, monkeypatch):
        # store_ppt writes images to the hardcoded relative folder data/images/.
        (tmp_path / "data" / "images").mkdir(parents=True)
        monkeypatch.chdir(tmp_path)
        yield

    def test_text_and_images_stored_in_separate_collections(self):
        with patch.object(dp, "extract_ppt_content",
                          return_value=("slide text", [self._fake_image()])), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]) as mock_chunk, \
             patch.object(dp, "get_image_embedding", return_value=[0.1] * 768) as mock_embed, \
             patch.object(dp, "store_in_vector_db") as mock_store:
            mock_store.side_effect = lambda **kw: ("col", ["id"])
            result = dp.store_ppt(path="deck.pptx", model_name="ppt_model")

        assert "completed" in result

        # Text collection.
        text_call = mock_store.call_args_list[0]
        assert text_call.kwargs["model_name"] == "ppt_model"
        assert text_call.kwargs["documents"] == ["chunk"]
        assert "embeddings" not in text_call.kwargs

        # Image collection, separate name, custom embeddings.
        image_call = mock_store.call_args_list[1]
        assert image_call.kwargs["model_name"] == "ppt_model_images"
        assert image_call.kwargs["embeddings"] == [[0.1] * 768]
        mock_embed.assert_called_once_with(b"png-bytes", mime_type="image/png")

        # The image bytes were persisted to disk so retrieval can reload them.
        import os
        saved = list(os.listdir("data/images"))
        assert len(saved) == 1 and saved[0].endswith(".png")

    def test_no_text_content_still_stores_images(self):
        with patch.object(dp, "extract_ppt_content", return_value=("   ", [self._fake_image()])), \
             patch.object(dp, "get_image_embedding", return_value=[0.1] * 768), \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])) as mock_store:
            dp.store_ppt(path="deck.pptx", model_name="ppt_model")

        # Only the image collection call happened.
        assert len(mock_store.call_args_list) == 1
        assert mock_store.call_args.kwargs["model_name"] == "ppt_model_images"

    def test_no_images_still_stores_text(self):
        with patch.object(dp, "extract_ppt_content", return_value=("some text", [])), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]), \
             patch.object(dp, "get_image_embedding") as mock_embed, \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])) as mock_store:
            dp.store_ppt(path="deck.pptx", model_name="ppt_model")

        mock_embed.assert_not_called()
        assert mock_store.call_args.kwargs["model_name"] == "ppt_model"

    def test_chunking_failure_aborts(self):
        with patch.object(dp, "extract_ppt_content", return_value=("text", [])), \
             patch.object(dp, "chunked_documents", return_value=None), \
             patch.object(dp, "store_in_vector_db") as mock_store:
            result = dp.store_ppt(path="deck.pptx", model_name="ppt_model")

        mock_store.assert_not_called()
        assert result is None

    def test_multiple_images_each_get_embedding_and_file(self):
        images = [self._fake_image(), self._fake_image()]
        with patch.object(dp, "extract_ppt_content", return_value=("text", images)), \
             patch.object(dp, "chunked_documents", return_value=["chunk"]), \
             patch.object(dp, "get_image_embedding",
                          side_effect=[[0.1] * 768, [0.2] * 768]) as mock_embed, \
             patch.object(dp, "store_in_vector_db", return_value=("col", ["id"])) as mock_store:
            dp.store_ppt(path="deck.pptx", model_name="ppt_model")

        assert mock_embed.call_count == 2
        image_call = mock_store.call_args_list[1]
        assert len(image_call.kwargs["documents"]) == 2
        assert image_call.kwargs["embeddings"] == [[0.1] * 768, [0.2] * 768]

        import os
        assert len(os.listdir("data/images")) == 2
